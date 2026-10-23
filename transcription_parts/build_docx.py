# Rebuilds "Full Kidase Transcription.docx" from the batch .md files in this folder.
import glob, os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(os.path.dirname(HERE), 'Full Kidase Transcription.docx')
ETH, LAT = 'Nyala', 'Times New Roman'

def set_font(run, name, size=None, bold=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = rpr.makeelement(qn('w:rFonts'), {}); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rf.set(qn(a), name)
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold

doc = Document()
st = doc.styles['Normal']; st.font.name = ETH; st.font.size = Pt(12)
h = doc.add_heading(level=0); set_font(h.add_run('ሥርዓተ ቅዳሴ — Full Kidase Transcription'), ETH)
p = doc.add_paragraph(); set_font(p.add_run('PDF pages 8–420 · ግዕዝ / አማርኛ / English'), ETH, 11)

# Liturgical speaker markers are bolded wherever they occur.
SPEAKER = re.compile(r'(ይ\.\s?[ሕዲካ]\.?|ይ\.\s?[ሕዲካ]\.?[ሕዲካ]?\.?|ካህን፦|ዲያቆን፦|ሕዝብ፦|Priest:|Deacon:|People:|Choir:)')
def add_text(para, text, font):
    for part in SPEAKER.split(text):
        if not part: continue
        set_font(para.add_run(part), font, 12, bool(SPEAKER.fullmatch(part)) or None)

LABELS = {'ግዕዝ': ETH, 'አማርኛ': ETH, 'English': LAT}
for f in sorted(glob.glob(os.path.join(HERE, 'p*.md'))):
    for line in open(f, encoding='utf-8').read().splitlines():
        line = line.rstrip()
        if not line or line == '---':
            continue
        if line.startswith('## '):
            hp = doc.add_heading(level=1); set_font(hp.add_run(line[3:]), ETH, None, True)
        elif line.startswith('### '):
            hp = doc.add_heading(level=2); set_font(hp.add_run(line[4:]), ETH, None, True)
        else:
            m = re.match(r'\*\*(ግዕዝ|አማርኛ|English):\*\*\s*(.*)', line)
            para = doc.add_paragraph(); para.paragraph_format.space_after = Pt(4)
            if m:
                set_font(para.add_run(m.group(1) + ': '), ETH, 12, True)
                add_text(para, m.group(2), LABELS[m.group(1)])
            else:
                add_text(para, line, ETH)
try:
    doc.save(OUT)
    print('saved', OUT)
except PermissionError:
    print('SKIPPED: docx is open in Word; will rebuild on next batch')
